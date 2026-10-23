"""Resume vs job description analysis (ATS-style score out of 100)."""
import io
import re
from collections import Counter

import pdfplumber
from docx import Document
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity

from skills_list import SKILLS

STOPWORDS = set("""
a about above after again all also am an and any are as at be because been before being below between both but by
can could did do does doing down during each few for from further had has have having he her here hers him his how
i if in into is it its just me more most my no nor not now of off on once only or other our out over own same she
should so some such than that the their them then there these they this those through to too under until up very
was we were what when where which while who whom why will with would you your yours
""".split())

GENERIC = set("""
experience work working team ability skills skill role company job candidate candidates required requirements
preferred responsibilities responsible years year strong good excellent knowledge understanding looking join
opportunity position apply including includes include etc plus must should able need needs using use used
well new get across based within across help support develop developing development manage managing
graduate engineer trainee fresher location salary benefits qualification qualifications eligibility
""".split())

SECTIONS = {
    "Contact info (email + phone)": None,  # handled separately
    "Education": r"\beducation\b|\bb\.?\s?tech\b|\bbachelor|\bdegree\b|\buniversity\b|\bcollege\b|\binstitute\b|\bcgpa\b",
    "Skills": r"\bskills?\b|\btechnologies\b|\btools\b",
    "Experience / Projects": r"\bexperience\b|\bprojects?\b|\binternships?\b|\bwork history\b",
}

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_RE = re.compile(r"(?:\+?\d{1,3}[\s-]?)?\d{10}|\+?\d[\d\s\-()]{8,}\d")


# ---------- reading files ----------
def extract_text(filename, data):
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    pages = None
    if ext == "pdf":
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = len(pdf.pages)
            text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        # some PDFs write bullet symbols as (cid:127); turn them into normal bullets
        text = re.sub(r"^\(cid:\d+\)\s*", "• ", text, flags=re.M)
        text = re.sub(r"\(cid:\d+\)", "", text)
    elif ext == "docx":
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        text = "\n".join(parts)
    elif ext == "txt":
        text = data.decode("utf-8", errors="ignore")
    else:
        raise ValueError("Please upload a PDF, DOCX or TXT file.")
    return text, pages


# ---------- keywords ----------
def _pattern(term):
    return re.compile(r"(?<![a-z0-9+#])" + re.escape(term) + r"(?:s|es)?(?![a-z0-9+#])")


def extract_keywords(jd, limit=25):
    low = jd.lower()
    counted = []
    for s in SKILLS:
        n = len(_pattern(s).findall(low))
        if n:
            counted.append((n, s))
    counted.sort(key=lambda x: -x[0])
    found = [s for _, s in counted]

    tokens = [w.strip(".-") for w in re.findall(r"[a-z][a-z+#.\-]{2,}", low)]
    counts = Counter(w for w in tokens if w not in STOPWORDS and w not in GENERIC and len(w) > 2)
    for w, c in counts.most_common():
        if len(found) >= limit:
            break
        if any(w == f or w in f.split() for f in found):
            continue
        if c >= 2 or len(found) < 8:
            found.append(w)
    return found[:limit]


# ---------- main analysis ----------
def analyze(filename, data, jd):
    text, pages = extract_text(filename, data)
    low = text.lower()
    result = {"text_ok": len(text.strip()) >= 50}
    if not result["text_ok"]:
        return result

    # 1) keywords (60 points: 45 keyword match + 15 text similarity)
    keywords = extract_keywords(jd)
    matched = [k for k in keywords if _pattern(k).search(low)]
    missing = [k for k in keywords if k not in matched]
    ratio = len(matched) / len(keywords) if keywords else 0
    try:
        vec = TfidfVectorizer(stop_words="english", ngram_range=(1, 2)).fit([low, jd.lower()])
        sim = float(cosine_similarity(vec.transform([low]), vec.transform([jd.lower()]))[0][0])
    except ValueError:
        sim = 0.0
    keyword_points = round(45 * ratio + 15 * min(sim / 0.5, 1.0), 1)

    # 2) sections (20 points, 5 each)
    section_checks = []
    for name, pat in SECTIONS.items():
        if pat is None:
            ok = bool(EMAIL_RE.search(text)) and bool(PHONE_RE.search(text))
        else:
            ok = bool(re.search(pat, low))
        section_checks.append({"name": name, "ok": ok})
    section_points = 5 * sum(c["ok"] for c in section_checks)

    # 3) formatting (20 points)
    words = len(re.findall(r"\w+", text))
    lines = [ln.strip() for ln in text.splitlines()]
    bullets = sum(1 for ln in lines if ln[:1] in "•-*▪●◦·" and len(ln) > 3)
    has_links = bool(re.search(r"linkedin\.com|github\.com", low))
    numbers = len(re.findall(r"\d+\s?%|\b\d{2,}\+?\b", text))

    length_pts = 8 if 250 <= words <= 900 else 4 if (150 <= words < 250 or 900 < words <= 1200) else 0
    pages_pts = 6 if (pages is None or pages <= 2) else 2 if pages == 3 else 0
    bullet_pts = 3 if bullets >= 3 else 0
    link_pts = 3 if has_links else 0
    format_points = length_pts + pages_pts + bullet_pts + link_pts
    format_checks = [
        {"name": f"Length is reasonable ({words} words)", "ok": length_pts == 8},
        {"name": "1-2 pages" + (f" ({pages} found)" if pages else ""), "ok": pages_pts == 6},
        {"name": "Uses bullet points", "ok": bullet_pts == 3},
        {"name": "Has LinkedIn or GitHub link", "ok": link_pts == 3},
    ]

    total = round(keyword_points + section_points + format_points)

    # tips
    tips = []
    if missing:
        tips.append("Add these missing keywords where they honestly apply to you: " + ", ".join(missing[:8]) + ".")
    for c in section_checks:
        if not c["ok"]:
            tips.append(f"Add a clear '{c['name'].split(' (')[0]}' section with a normal heading.")
    if words < 250:
        tips.append("Your resume looks short. Add project details and what you achieved.")
    if words > 900:
        tips.append("Your resume is long. Keep it to 1 page (fresher) or 2 pages.")
    if pages and pages > 2:
        tips.append("Reduce your resume to 2 pages or fewer.")
    if bullets < 3:
        tips.append("Use bullet points (•) for projects and experience so they are easy to scan.")
    if not has_links:
        tips.append("Add your LinkedIn and GitHub links.")
    if numbers < 3:
        tips.append("Add numbers to show results (for example: 'reduced response time by 30%').")
    if not tips:
        tips.append("Great job! Re-check that every keyword you list is something you can talk about in an interview.")

    result.update({
        "score": total,
        "band": "Strong match" if total >= 75 else "Good, can improve" if total >= 50 else "Needs work",
        "keyword_points": keyword_points, "section_points": section_points, "format_points": format_points,
        "matched": matched, "missing": missing,
        "section_checks": section_checks, "format_checks": format_checks,
        "tips": tips, "words": words, "similarity": round(sim * 100),
    })
    return result
